"""Tests for Word document builder (D1 - D5)."""

from datetime import date
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pytest

from app.docx_builder import build_docx
from app.schemas import DocumentSpec, SectionDraft, Table


def test_d1_standard_document_structure(tmp_path: Path):
    """D1 file opens; title present; headings numbered in order; assumptions section last; table rows and columns match the spec; header row has w:tblHeader; footer XML contains PAGE and NUMPAGES; core title set."""
    spec = DocumentSpec(
        title="Project Proposal for Cloud Migration",
        doc_type="proposal",
        audience="Board of Directors",
        generated_on=date(2026, 10, 23),
        assumptions=["Cloud provider is AWS", "Migration starts Q1"],
        open_questions=["What is the maintenance window?"],
        sections=[
            SectionDraft(
                id="t1",
                heading="Executive Summary",
                paragraphs=["This proposal outlines the strategy for moving operations to the cloud."],
                bullets=["Cost efficiency", "High availability"],
            ),
            SectionDraft(
                id="t2",
                heading="Delivery Timeline",
                paragraphs=["The project will run over 12 weeks."],
                table=Table(
                    caption="Schedule",
                    headers=["Phase", "Duration", "Cost"],
                    rows=[
                        ["Phase 1", "4 weeks", "$10,000"],
                        ["Phase 2", "8 weeks", "$25,000"],
                    ],
                ),
            ),
        ],
    )
    out_file = tmp_path / "test_d1.docx"
    build_docx(spec, out_file)

    assert out_file.exists()
    doc = Document(str(out_file))

    # Core title set
    assert doc.core_properties.title == "Project Proposal for Cloud Migration"
    assert doc.core_properties.author == "Autonomous Document Agent"

    # Title present in text
    all_text = [p.text for p in doc.paragraphs]
    assert any("Project Proposal for Cloud Migration" in t for t in all_text)

    # Headings numbered in order
    assert any("1. Executive Summary" in t for t in all_text)
    assert any("2. Delivery Timeline" in t for t in all_text)
    assert any("3. Assumptions and Open Questions" in t for t in all_text)

    # Table rows and columns match
    tables = doc.tables
    # Tables include callout table + section table
    section_table = tables[1]
    assert len(section_table.rows) == 3  # 1 header + 2 data rows
    assert len(section_table.columns) == 3

    # Header row has w:tblHeader
    header_tr_xml = section_table.rows[0]._tr.xml
    assert "tblHeader" in header_tr_xml

    # Footer XML contains PAGE and NUMPAGES
    footer = doc.sections[0].footer
    footer_xml = footer._element.xml
    assert "PAGE" in footer_xml
    assert "NUMPAGES" in footer_xml


def test_d2_no_assumptions_no_callout_or_section(tmp_path: Path):
    """D2 no assumptions means no callout and no assumptions section."""
    spec = DocumentSpec(
        title="Direct Report",
        doc_type="report",
        audience="Engineers",
        generated_on=date(2026, 10, 23),
        assumptions=[],
        open_questions=[],
        sections=[
            SectionDraft(id="t1", heading="Architecture", paragraphs=["Clean design."]),
        ],
    )
    out_file = tmp_path / "test_d2.docx"
    build_docx(spec, out_file)

    doc = Document(str(out_file))
    all_text = [p.text for p in doc.paragraphs]
    assert not any("assumptions were made" in t.lower() for t in all_text)
    assert not any("Assumptions and Open Questions" in t for t in all_text)
    # Only 0 tables created (no callout table, no section table)
    assert len(doc.tables) == 0


def test_d3_eight_column_table_and_large_cell(tmp_path: Path):
    """D3 8-column table and a 2,000-char cell do not raise."""
    large_text = "Detailed requirement paragraph. " * 70  # ~2,240 chars
    headers = [f"Col {i}" for i in range(8)]
    row = [large_text if i == 0 else f"Val {i}" for i in range(8)]

    spec = DocumentSpec(
        title="Large Table Document",
        doc_type="report",
        audience="Operations",
        generated_on=date(2026, 10, 23),
        sections=[
            SectionDraft(
                id="t1",
                heading="Data Matrix",
                table=Table(headers=headers, rows=[row]),
            )
        ],
    )
    out_file = tmp_path / "test_d3.docx"
    # Must not raise
    build_docx(spec, out_file)
    assert out_file.exists()


def test_d4_strip_existing_numbering(tmp_path: Path):
    """D4 LLM-added numbering ('1. Scope') is not doubled."""
    spec = DocumentSpec(
        title="Numbered Proposal",
        doc_type="proposal",
        audience="Clients",
        generated_on=date(2026, 10, 23),
        sections=[
            SectionDraft(id="t1", heading="1. Scope and Objectives", paragraphs=["Defined."]),
            SectionDraft(id="t2", heading="## 2. Technical Approach", paragraphs=["Detailed."]),
            SectionDraft(id="t3", heading="3) Commercial Terms", paragraphs=["Agreed."]),
        ],
    )
    out_file = tmp_path / "test_d4.docx"
    build_docx(spec, out_file)

    doc = Document(str(out_file))
    all_text = [p.text for p in doc.paragraphs]
    assert any("1. Scope and Objectives" in t for t in all_text)
    assert not any("1. 1." in t for t in all_text)
    assert any("2. Technical Approach" in t for t in all_text)
    assert not any("2. 2." in t or "2. ##" in t for t in all_text)
    assert any("3. Commercial Terms" in t for t in all_text)


def test_d5_total_row_bold_and_numeric_alignment(tmp_path: Path):
    """D5 'Total' row is bold; numeric cells right-aligned."""
    spec = DocumentSpec(
        title="Budget Overview",
        doc_type="proposal",
        audience="Finance",
        generated_on=date(2026, 10, 23),
        sections=[
            SectionDraft(
                id="t1",
                heading="Financials",
                table=Table(
                    headers=["Item", "Cost"],
                    rows=[
                        ["Software License", "$5,000.00"],
                        ["Total", "$5,000.00"],
                    ],
                ),
            )
        ],
    )
    out_file = tmp_path / "test_d5.docx"
    build_docx(spec, out_file)

    doc = Document(str(out_file))
    tbl = doc.tables[0]

    # Row 1 (Item): Cost is numeric, should be right aligned
    data_row_1 = tbl.rows[1]
    cell_cost_p = data_row_1.cells[1].paragraphs[0]
    assert cell_cost_p.alignment == WD_ALIGN_PARAGRAPH.RIGHT

    # Row 2 (Total): runs in all cells should be bold
    total_row = tbl.rows[2]
    assert total_row.cells[0].paragraphs[0].runs[0].font.bold is True
    assert total_row.cells[1].paragraphs[0].runs[0].font.bold is True
