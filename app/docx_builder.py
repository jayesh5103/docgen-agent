"""High-quality Word (.docx) document generator using python-docx."""

from datetime import date
from pathlib import Path
import re
from typing import Any
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

from app.schemas import DocumentSpec, Table

# Styling constants
COLOR_PRIMARY = RGBColor(0x1F, 0x38, 0x64)     # Dark Blue
COLOR_BODY = RGBColor(0x26, 0x26, 0x26)        # Near Black / Dark Charcoal
COLOR_MUTED = RGBColor(0x6B, 0x72, 0x80)       # Gray
HEX_PRIMARY = "1F3864"
HEX_ZEBRA = "F8FAFC"
HEX_BORDER = "D1D5DB"
HEX_CALLOUT_BG = "F0F4F8"


def _set_cell_background(cell: Any, hex_color: str) -> None:
    """Set background shading for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell: Any, top: int = 120, bottom: int = 120, left: int = 160, right: int = 160) -> None:
    """Set internal cell padding (in dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def _set_table_borders(table: Any) -> None:
    """Apply subtle light gray borders to table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{HEX_PRIMARY}"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


def _is_numeric(text: str) -> bool:
    """Heuristic to check if cell content is numeric, monetary, or percentage."""
    cleaned = text.strip()
    if not cleaned:
        return False
    # Strip currency symbols and commas
    stripped = re.sub(r"^[\$€£₹A-Za-z\s]+", "", cleaned)
    stripped = stripped.replace(",", "").replace("%", "").strip()
    try:
        float(stripped)
        return True
    except ValueError:
        return False


def _clean_heading(heading: str) -> str:
    """Strip pre-existing numbers, markdown hashes, or bullets from heading."""
    cleaned = heading.strip()
    pattern = r"^(\#+|\d+[\.\)]|\*|\-|\:)\s*"
    while re.match(pattern, cleaned):
        cleaned = re.sub(pattern, "", cleaned).strip()
    return cleaned


def _add_page_number_fields(paragraph: Any) -> None:
    """Insert Word PAGE and NUMPAGES dynamic fields into paragraph."""
    p_element = paragraph._p

    # "Page "
    run1 = paragraph.add_run("Page ")
    run1.font.name = "Calibri"
    run1.font.size = Pt(9)
    run1.font.color.rgb = COLOR_MUTED

    # PAGE field
    fld_page = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls("w"))
    p_element.append(fld_page)

    # " of "
    run2 = paragraph.add_run(" of ")
    run2.font.name = "Calibri"
    run2.font.size = Pt(9)
    run2.font.color.rgb = COLOR_MUTED

    # NUMPAGES field
    fld_numpages = parse_xml(r'<w:fldSimple %s w:instr="NUMPAGES"/>' % nsdecls("w"))
    p_element.append(fld_numpages)


def build_docx(spec: DocumentSpec, path: Path | str) -> Path:
    """Build a styled, professional Word document from a DocumentSpec."""
    target_path = Path(path)
    target_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # 1. Page Setup: A4, 2.2 cm margins
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # 2. Font configuration (Calibri 11pt default, strip theme fonts)
    for style_name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            # Strip theme fonts in XML
            rPr = style.element.get_or_add_rPr()
            for child in list(rPr):
                if child.tag.endswith("rFonts"):
                    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
                        attr_q = f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attr}"
                        if attr_q in child.attrib:
                            del child.attrib[attr_q]

    normal_style = doc.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_BODY

    # 3. Header & Footer
    header_para = section.header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hdr_run = header_para.add_run(spec.doc_type.replace("_", " ").title())
    hdr_run.font.name = "Calibri"
    hdr_run.font.size = Pt(8.5)
    hdr_run.font.color.rgb = COLOR_MUTED

    footer_para = section.footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_page_number_fields(footer_para)

    # 4. Title block
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    t_run = title_p.add_run(spec.title)
    t_run.font.name = "Calibri"
    t_run.font.size = Pt(22)
    t_run.font.bold = True
    t_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(16)
    sub_text = (
        f"{spec.doc_type.replace('_', ' ').title()} · "
        f"{spec.generated_on.isoformat()} · "
        "Draft by autonomous agent"
    )
    s_run = sub_p.add_run(sub_text)
    s_run.font.name = "Calibri"
    s_run.font.size = Pt(9.5)
    s_run.font.color.rgb = COLOR_MUTED

    # 5. Shaded callout under title if assumptions exist
    has_assumptions = bool(spec.assumptions)
    has_open_questions = bool(spec.open_questions)
    final_section_needed = has_assumptions or has_open_questions
    final_section_num = len(spec.sections) + 1 if final_section_needed else None

    if has_assumptions:
        callout_tbl = doc.add_table(rows=1, cols=1)
        callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        callout_cell = callout_tbl.cell(0, 0)
        callout_cell.width = Cm(16.6)
        _set_cell_background(callout_cell, HEX_CALLOUT_BG)
        _set_cell_margins(callout_cell, top=140, bottom=140, left=200, right=200)

        c_para = callout_cell.paragraphs[0]
        c_para.paragraph_format.space_before = Pt(0)
        c_para.paragraph_format.space_after = Pt(0)
        c_run = c_para.add_run(
            f"ℹ️ {len(spec.assumptions)} key assumptions were made for this document; "
            f"see Section {final_section_num} for details."
        )
        c_run.font.name = "Calibri"
        c_run.font.size = Pt(9.5)
        c_run.font.italic = True
        c_run.font.color.rgb = COLOR_PRIMARY

        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(6)
        spacer.paragraph_format.space_after = Pt(6)

    # 6. Sections in Plan Order
    for sec_idx, draft in enumerate(spec.sections, start=1):
        clean_title = _clean_heading(draft.heading)
        h_para = doc.add_paragraph()
        h_para.paragraph_format.space_before = Pt(16)
        h_para.paragraph_format.space_after = Pt(6)
        h_para.paragraph_format.keep_with_next = True

        h_run = h_para.add_run(f"{sec_idx}. {clean_title}")
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(14)
        h_run.font.bold = True
        h_run.font.color.rgb = COLOR_PRIMARY

        # Paragraphs
        for para_text in draft.paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            p_run = p.add_run(para_text)
            p_run.font.name = "Calibri"
            p_run.font.size = Pt(11)

        # Bullets
        for b_text in draft.bullets:
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_before = Pt(0)
            bp.paragraph_format.space_after = Pt(3)
            b_run = bp.add_run(b_text)
            b_run.font.name = "Calibri"
            b_run.font.size = Pt(10.5)

        # Table
        if draft.table and draft.table.headers:
            t_spec: Table = draft.table
            if t_spec.caption:
                cap_para = doc.add_paragraph()
                cap_para.paragraph_format.space_before = Pt(8)
                cap_para.paragraph_format.space_after = Pt(2)
                cap_run = cap_para.add_run(f"Table {sec_idx}: {t_spec.caption}")
                cap_run.font.name = "Calibri"
                cap_run.font.size = Pt(9.5)
                cap_run.font.italic = True
                cap_run.font.color.rgb = COLOR_MUTED

            num_cols = len(t_spec.headers)
            num_rows = len(t_spec.rows) + 1
            tbl = doc.add_table(rows=num_rows, cols=num_cols)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            _set_table_borders(tbl)

            # Header row configuration
            hdr_row = tbl.rows[0]
            # w:tblHeader repeating across pages
            trPr = hdr_row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            # Compute column width proportions based on content
            total_avail_cm = 16.6
            col_weights = []
            for col_idx in range(num_cols):
                max_len = len(str(t_spec.headers[col_idx]))
                for r in t_spec.rows:
                    if col_idx < len(r):
                        max_len = max(max_len, min(len(str(r[col_idx])), 80))
                col_weights.append(max(max_len, 4))
            sum_weights = sum(col_weights) or 1
            col_widths = [Cm(max(1.5, (w / sum_weights) * total_avail_cm)) for w in col_weights]

            for c_idx, h_text in enumerate(t_spec.headers):
                cell = hdr_row.cells[c_idx]
                cell.width = col_widths[c_idx]
                _set_cell_background(cell, HEX_PRIMARY)
                _set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                hrun = p.add_run(h_text)
                hrun.font.name = "Calibri"
                hrun.font.size = Pt(9.5)
                hrun.font.bold = True
                hrun.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Data rows
            for r_idx, row_data in enumerate(t_spec.rows, start=1):
                d_row = tbl.rows[r_idx]
                is_zebra = (r_idx % 2 == 0)
                first_cell_text = str(row_data[0]).strip().lower() if row_data else ""
                is_total_row = first_cell_text.startswith("total") or first_cell_text.startswith("subtotal")

                for c_idx in range(num_cols):
                    cell = d_row.cells[c_idx]
                    cell.width = col_widths[c_idx]
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    _set_cell_margins(cell, top=100, bottom=100, left=140, right=140)

                    if is_zebra and not is_total_row:
                        _set_cell_background(cell, HEX_ZEBRA)

                    val = str(row_data[c_idx]) if c_idx < len(row_data) else ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)

                    if _is_numeric(val) and not is_total_row and c_idx > 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    d_run = p.add_run(val)
                    d_run.font.name = "Calibri"
                    d_run.font.size = Pt(9.5)
                    if is_total_row:
                        d_run.font.bold = True

            spacer_table = doc.add_paragraph()
            spacer_table.paragraph_format.space_before = Pt(4)
            spacer_table.paragraph_format.space_after = Pt(4)

    # 7. Final section: Assumptions and Open Questions (from plan, not LLM)
    if final_section_needed:
        f_para = doc.add_paragraph()
        f_para.paragraph_format.space_before = Pt(16)
        f_para.paragraph_format.space_after = Pt(6)
        f_para.paragraph_format.keep_with_next = True

        f_run = f_para.add_run(f"{final_section_num}. Assumptions and Open Questions")
        f_run.font.name = "Calibri"
        f_run.font.size = Pt(14)
        f_run.font.bold = True
        f_run.font.color.rgb = COLOR_PRIMARY

        if spec.assumptions:
            intro_p = doc.add_paragraph()
            intro_p.paragraph_format.space_before = Pt(0)
            intro_p.paragraph_format.space_after = Pt(4)
            i_run = intro_p.add_run("The following assumptions were made to prepare this document:")
            i_run.font.bold = True
            i_run.font.size = Pt(10.5)

            for asm in spec.assumptions:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(3)
                r = bp.add_run(asm)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)

        if spec.open_questions:
            oq_p = doc.add_paragraph()
            oq_p.paragraph_format.space_before = Pt(6)
            oq_p.paragraph_format.space_after = Pt(4)
            oq_run = oq_p.add_run("Open questions requiring stakeholder confirmation:")
            oq_run.font.bold = True
            oq_run.font.size = Pt(10.5)

            for q in spec.open_questions:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_before = Pt(0)
                bp.paragraph_format.space_after = Pt(3)
                r = bp.add_run(q)
                r.font.name = "Calibri"
                r.font.size = Pt(10.5)

    # 8. Core properties
    doc.core_properties.title = spec.title
    doc.core_properties.author = "Autonomous Document Agent"
    doc.core_properties.subject = spec.doc_type

    doc.save(str(target_path))
    return target_path
